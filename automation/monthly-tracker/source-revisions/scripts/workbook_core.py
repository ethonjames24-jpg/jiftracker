#!/usr/bin/env python3
"""Read-only reconciliation of the six pinned official XLSX snapshots.

Evaluates the limited, explicitly checked arithmetic/SUM/ABS formula grammar
used by these workbooks. Does not invoke Excel, write XLSX, or execute macros.
"""
from pathlib import Path
from decimal import Decimal, getcontext, ROUND_HALF_UP
from collections import Counter
import ast
import csv
import hashlib
import json
import re
import sys
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries, get_column_letter

getcontext().prec = 35
ROOT = Path(__file__).resolve().parent
WORK = ROOT / 'JIF-April-June-Source-Reconciliation-2026-10-28'
SOURCE = WORK / 'sources'
Q = Decimal('0.001')
EPS = Decimal('0.0000001')
LABELS = ['Tax Revenue', 'Non-Tax Revenue', 'Total Revenue & Grants',
          'Primary Balance', 'Fiscal Balance', 'Compensation of Employees',
          'Capital Expenditure', 'Interest']
ROWS = dict(zip(LABELS, [10, 11, 9, 44, 29, 19, 26, 22]))
PAIRS = {
 '2026-04': ('Copy-of-CGO-April-2026.xlsx', 'Copy-of-CGO-April-2026-1.xlsx', 'G', 'H'),
 '2026-05': ('Copy-of-CGO-May-2026-Revised.xlsx', 'Copy-of-CGO-May-2026_.xlsx', 'I', 'J'),
 '2026-06': ('Copy-of-CGO-June-2026.xlsx', 'Copy-of-CGO-June-2026-1.xlsx', 'J', 'K'),
}

def decimal(value):
    if value is None:
        return Decimal(0)
    if isinstance(value, bool):
        raise ValueError('Boolean cannot be a fiscal amount')
    return Decimal(str(value))

class Evaluator:
    TOKEN = re.compile(r'\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?')

    def __init__(self, sheet):
        self.sheet, self.cache, self.stack, self.dependencies = sheet, {}, set(), {}

    def cell(self, address):
        address = address.replace('$', '')
        if address in self.cache:
            return self.cache[address]
        if address in self.stack:
            raise ValueError(f'Circular formula: {address}')
        self.stack.add(address)
        value = self.sheet[address].value
        self.dependencies[address] = set()
        if isinstance(value, str) and value.startswith('='):
            if '!' in value or '[' in value:
                raise ValueError(f'External or cross-sheet reference: {address}')
            def token(match):
                t = match[0].replace('$', '')
                if ':' in t:
                    a,b=t.split(':');return f'R("{a}","{b}")'
                return f'C("{t}")'
            expression = self.TOKEN.sub(token, value[1:])
            parsed = ast.parse(expression, mode='eval')
            result = self.node(parsed.body, address)
        elif value is None or isinstance(value, (int, float, Decimal)):
            result = decimal(value)
        else:
            result = value
        self.stack.remove(address)
        self.cache[address] = result
        return result

    def node(self, node, owner):
        if isinstance(node, ast.Constant):
            return decimal(node.value) if isinstance(node.value,(int,float)) else node.value
        if isinstance(node, ast.UnaryOp) and isinstance(node.op,(ast.UAdd,ast.USub)):
            value=self.node(node.operand,owner)
            return value if isinstance(node.op,ast.UAdd) else -value
        if isinstance(node, ast.BinOp):
            a,b=self.node(node.left,owner),self.node(node.right,owner)
            if not isinstance(a,Decimal) or not isinstance(b,Decimal):
                raise ValueError(f'Non-numeric arithmetic at {owner}')
            if isinstance(node.op,ast.Add):return a+b
            if isinstance(node.op,ast.Sub):return a-b
            if isinstance(node.op,ast.Mult):return a*b
            if isinstance(node.op,ast.Div):return a/b
        if isinstance(node,ast.Call) and isinstance(node.func,ast.Name):
            args=[self.node(x,owner) for x in node.args]
            if node.func.id=='C' and len(args)==1:
                target=args[0];self.dependencies[owner].add(target)
                result=self.cell(target)
                self.dependencies[owner].update(self.dependencies[target])
                return result
            if node.func.id=='R' and len(args)==2:
                c1,r1,c2,r2=range_boundaries(':'.join(args));values=[]
                for row in range(r1,r2+1):
                    for col in range(c1,c2+1):
                        target=f'{get_column_letter(col)}{row}'
                        self.dependencies[owner].add(target);values.append(self.cell(target))
                        self.dependencies[owner].update(self.dependencies[target])
                return values
            if node.func.id=='SUM':
                values=[v for a in args for v in (a if isinstance(a,list) else [a])]
                return sum((v for v in values if isinstance(v,Decimal)),Decimal(0))
            if node.func.id=='ABS' and len(args)==1:return abs(args[0])
        raise ValueError(f'Unsupported formula grammar at {owner}: {ast.dump(node)}')

def status(label, gap):
    if label in ('Compensation of Employees','Interest'):
        return 'On Track' if gap<=Decimal('.03') else 'Watch' if gap<=Decimal('.08') else 'Under Pressure'
    on,watch=(Decimal('-.10'),Decimal('-.20')) if label=='Capital Expenditure' else (Decimal('-.03'),Decimal('-.08'))
    return 'On Track' if gap>=on else 'Watch' if gap>=watch else 'Under Pressure'

def overall(counts):
    u,w=counts['Under Pressure'],counts['Watch']
    return 'Under Pressure' if u>=3 else 'Watch' if u>=1 or w>=3 else 'On Track'

def published_amount(text):
    return Decimal(re.sub(r'[^0-9.\-]', '', text))

def serialize(value):
    if isinstance(value,Decimal):return str(value)
    if isinstance(value,set):return sorted(value)
    raise TypeError(type(value))
