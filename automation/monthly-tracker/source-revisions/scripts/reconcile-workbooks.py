#!/usr/bin/env python3
"""Read-only workbook comparison; no Excel execution, macros, source edits or network."""
import argparse,json,hashlib,zipfile,re
from pathlib import Path
from decimal import Decimal
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from workbook_core import Evaluator,ROWS,LABELS,status,serialize,EPS
MONTHS=['Jan','Feb','Mar','Apr','May','Jun','Jul','Aug','Sep','Oct','Nov','Dec']
def extract(path,month):
 with zipfile.ZipFile(path) as z:
  assert len(z.infolist())<=10000 and sum(x.file_size for x in z.infolist())<=50000000,'ZIP_EXPANSION_LIMIT'
 forms=load_workbook(path,data_only=False,keep_links=False)['A'];cached=load_workbook(path,data_only=True,keep_links=False)['A']
 year,mm=map(int,month.split('-'));fy=year-(mm<4);assert str(forms['B3'].value).strip()==f'FY {fy}/{str(fy+1)[-2:]}','FISCAL_YEAR_MISMATCH'
 assert 'millions of Jamaica dollars' in str(forms['B4'].value),'UNITS_MISMATCH'
 period='Apr' if mm==4 else 'Apr - '+MONTHS[mm-1]
 find=lambda kind:[c for c in range(1,cached.max_column+1) if str(cached.cell(6,c).value).strip()==kind and str(cached.cell(7,c).value).strip()==period]
 actual,budget=find('Prov.'),find('Budget');assert len(actual)==len(budget)==1,'YTD_COLUMN_AMBIGUITY'
 ev=Evaluator(forms);result={}
 for label,row in ROWS.items():
  found=next((str(forms.cell(row,c).value).strip() for c in range(2,5) if forms.cell(row,c).value is not None),'')
  expected={'Total Revenue & Grants':'Revenue & Grants','Primary Balance':'Primary Balance (Surplus + / Deficit -)','Fiscal Balance':'Fiscal Balance (Surplus + / Deficit -)'}.get(label,label);assert found==expected,'KPI_ROW_MISMATCH'
  a,b=f'{get_column_letter(actual[0])}{row}',f'{get_column_letter(budget[0])}{row}';av,bv=ev.cell(a),ev.cell(b);assert bv!=0,'ZERO_BASELINE_REVIEW_REQUIRED'
  for cell,value in [(a,av),(b,bv)]:assert abs(value-Decimal(str(cached[cell].value)))<=EPS,'CACHE_RECALCULATION_MISMATCH'
  gap=(av-bv)/abs(bv);result[label]={'actual':av,'budget':bv,'gap_pct':gap*100,'status':status(label,gap),'actual_cell':a,'budget_cell':b}
 return forms,result

def reconcile(old,new,month):
 a,ar=extract(old,month);b,br=extract(new,month)
 changes=[{'cell':a.cell(r,c).coordinate,'original':str(a.cell(r,c).value),'candidate':str(b.cell(r,c).value)} for r in range(1,max(a.max_row,b.max_row)+1) for c in range(1,max(a.max_column,b.max_column)+1) if a.cell(r,c).value!=b.cell(r,c).value]
 records=[{'kpi':k,'original':ar[k],'candidate':br[k],'amount_changed':any(ar[k][f]!=br[k][f] for f in ['actual','budget']),'status_changed':ar[k]['status']!=br[k]['status']} for k in LABELS]
 return {'month_sort':month,'units':'J$ millions','original_sha256':hashlib.sha256(old.read_bytes()).hexdigest(),'candidate_sha256':hashlib.sha256(new.read_bytes()).hexdigest(),'literal_changes':changes,'kpis':records,'amount_changed_rows':sum(r['amount_changed'] for r in records),'status_changed_rows':sum(r['status_changed'] for r in records),'cache_input_checks':32,'engine':'Restricted arithmetic/SUM/ABS evaluator; not native Excel certification; comparison scope worksheet A and its supported dependencies','financial_adoption_authorized':False,'public_writes':0}
if __name__=='__main__':
 p=argparse.ArgumentParser();p.add_argument('month');p.add_argument('original',type=Path);p.add_argument('candidate',type=Path);args=p.parse_args();assert re.fullmatch(r'\d{4}-(0[1-9]|1[0-2])',args.month);print(json.dumps(reconcile(args.original,args.candidate,args.month),default=serialize,indent=2))
